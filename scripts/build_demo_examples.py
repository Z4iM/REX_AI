"""Build explicitly fictional TA fixtures. Run with a Python containing python-docx and PyMuPDF."""
from pathlib import Path
from datetime import date, timedelta
from html import escape
import json
import pymupdf as fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'examples' / 'ta_demo'
MARKER = 'REX_SYNTHETIC_DEMO'
NOTICE = MARKER + ' | Fictional TA example. No real academy offering, instructor, accreditation or tender.'

COURSES = [
 dict(course_id='DEMO-PY-101', title='Python Foundations Demo', description='A fictional practical introduction to Python programming using variables, conditions, loops, functions and CSV files. Learners build a small expense tracker using invented transactions.', level='Beginner', training_format='Online', price_egp=2400, duration_hours=24, department='Software Development', instructor_name='Nour Demo Instructor', instructor_bio='Fictional instructor created solely for demonstration.', schedule='Demo cohort: 8 sessions of 3 hours, Sunday and Tuesday 18:00-21:00 Cairo time. Dates require confirmation.', offers='Demo offer: 10 percent off for a fictional group of five; this is not a real promotion.', target_audience='Beginners who want to automate simple office tasks.', career_opportunities=['Foundation for junior Python development', 'Office task automation'], admission_requirements=['Access to a laptop with Python 3.11 or later', 'Ability to manage local files'], prerequisites=['Basic computer use; no programming experience required'], learning_outcomes=['Write conditions, loops and reusable functions', 'Read CSV files and validate numeric values', 'Build and test an expense tracker'], faqs=[{'question':'Is programming experience required?', 'answer':'No. Basic computer use is enough for this fictional beginner course.'},{'question':'Is a job guaranteed?', 'answer':'No employment or accreditation promise is made.'}], modules=[('Getting started', 'Install Python and run print statements. Store an expense amount in a variable. Distinguish strings from numbers and convert input with float. Exercise: calculate a subtotal and 14 percent tax for invented items.'),('Decisions and repetition', 'Use if to reject negative amounts, for to total a list and while to repeat a menu. Exercise: process [120, 80, 50] and verify the total is 250. Explain how an empty list should behave.'),('Functions and files', 'Write add_expense(amount, category). Use csv.DictReader and handle missing or malformed amounts with a clear error. Exercise: load three fictional rows and skip an invalid row while reporting its line number.'),('Expense tracker project', 'Create a menu to add, list and total expenses. Keep data in a local CSV file. Submit the script, a README and tests for empty input, negative amounts and nonnumeric input. Assess correctness 50 percent, validation 30 percent and clarity 20 percent.')]),
 dict(course_id='DEMO-DA-201', title='Data Analysis with Python Demo', description='A fictional applied course in cleaning tabular data, calculating business metrics and communicating findings with Python and pandas. All exercises use invented sales data.', level='Intermediate', training_format='Hybrid', price_egp=4800, duration_hours=36, department='Data and AI', instructor_name='Salma Demo Instructor', instructor_bio='Fictional analyst instructor for demonstration only.', schedule='Demo cohort: 12 sessions of 3 hours, Monday and Wednesday 18:00-21:00 Cairo time; dates unconfirmed.', offers='No discount is specified for this fictional course.', target_audience='Learners comfortable with Python functions and CSV files.', career_opportunities=['Foundation for junior data analysis', 'Reporting and operations analytics'], admission_requirements=['Laptop with Python and pandas', 'Complete a basic Python exercise'], prerequisites=['Python variables, loops and functions', 'Basic spreadsheet arithmetic'], learning_outcomes=['Clean missing and duplicate values', 'Calculate grouped sales metrics', 'Explain a chart without making causal claims'], faqs=[{'question':'Can a complete beginner join?', 'answer':'Python foundations are required before this intermediate example.'}], modules=[('Inspect and clean', 'Read sales.csv with pandas. Check column types and missing values. Remove exact duplicate rows and document why an absent quantity must be checked before filling.'),('Calculate metrics', 'Calculate revenue as quantity multiplied by unit_price. Group revenue by category. Verify two units at 150 produce 300 and never sum unit prices as if they were revenue.'),('Visualize and communicate', 'Create a bar chart by category with labeled axes and currency. Explain sample size and missing data. State that a correlation does not demonstrate a causal relationship.'),('Analysis project', 'Submit a reproducible notebook with input checks, cleaned output, three findings and one limitation. Assess data quality 35 percent, calculations 35 percent and communication 30 percent.')]),
 dict(course_id='DEMO-CL-301', title='Cloud Operations Demo', description='A fictional advanced course in secure cloud operations, service monitoring, incident response and recovery planning. Exercises run locally or use documented examples; no real deployment or paid cloud account is required.', level='Advanced', training_format='On-site', price_egp=7200, duration_hours=40, department='Cloud and Infrastructure', instructor_name='Omar Demo Instructor', instructor_bio='Fictional infrastructure instructor; no real certifications claimed.', schedule='Demo cohort: 10 sessions of 4 hours on Saturdays; venue and dates require confirmation.', offers='No active promotion is specified.', target_audience='Developers and administrators with Linux and networking experience.', career_opportunities=['Cloud operations practice', 'Site reliability foundations'], admission_requirements=['Laptop supporting containers', 'A local container runtime'], prerequisites=['Linux command line', 'HTTP, DNS and basic networking', 'Basic containers'], learning_outcomes=['Define least privilege access', 'Design service health checks and alerts', 'Test backup restoration and incident steps'], faqs=[{'question':'Is a paid cloud account required?', 'answer':'No. The demonstration uses local containers and authored architecture examples.'}], modules=[('Access and boundaries', 'Map service identities to required operations. Remove wildcard permissions and explain least privilege. Exercise: draft a read-only access policy without embedding any secret.'),('Monitoring', 'Define availability and latency measures. Use a local health endpoint and simulate a failed dependency. Explain the difference between liveness and readiness.'),('Incident handling', 'Document detection, triage, mitigation and follow-up. Exercise: investigate an intentionally stopped local container using logs and a timeline, without deleting persistent volumes.'),('Recovery project', 'Back up a fictional configuration directory and restore it to a new location. Record recovery time and verify checksums. Assess safety 40 percent, observability 30 percent and recovery evidence 30 percent.')]),
 dict(course_id='DEMO-AR-101', title='مقدمة في المهارات الرقمية التجريبية', description='دورة خيالية للعرض أمام المعيد عن تنظيم الملفات واستخدام الجداول وحماية الحسابات وكتابة رسائل واضحة. جميع الأشخاص والبيانات أمثلة تعليمية غير حقيقية.', level='Beginner', training_format='Online', price_egp=1800, duration_hours=18, department='Digital Skills', instructor_name='ليلى مدربة تجريبية', instructor_bio='شخصية خيالية لأغراض العرض فقط.', schedule='جدول تجريبي: ست جلسات، كل جلسة ثلاث ساعات. المواعيد غير مؤكدة.', offers='لا توجد عروض حقيقية؛ هذه بيانات تجريبية.', target_audience='مبتدئون في استخدام الأدوات الرقمية.', career_opportunities=['تحسين المهارات المكتبية الأساسية'], admission_requirements=['جهاز حاسوب واتصال بالإنترنت'], prerequisites=['القدرة على تشغيل الحاسوب'], learning_outcomes=['تنظيم الملفات وتسميتها بوضوح', 'حساب مجموع باستخدام جدول بيانات', 'تمييز الرسائل الاحتيالية'], faqs=[{'question':'هل هذه دورة حقيقية؟','answer':'لا. هذه دورة خيالية لاختبار النظام أمام المعيد.'}], modules=[('تنظيم الملفات','أنشئ مجلداً للمشروع ومجلداً للنسخ الاحتياطية. استخدم أسماء واضحة دون بيانات شخصية. تمرين: رتب خمسة ملفات تجريبية حسب الموضوع.'),('الجداول الأساسية','أدخل ثلاث قيم تجريبية واحسب مجموعها باستخدام صيغة. تحقق من أن مجموع ١٠ و٢٠ و٣٠ يساوي ٦٠، ثم راجع نوع البيانات.'),('أمان الحسابات','استخدم كلمة مرور فريدة والتحقق بخطوتين. لا تضع كلمات مرور حقيقية في ملفات التدريب. تمرين: حدد علامات رسالة احتيالية وهمية.'),('المشروع النهائي','أنشئ جدول نفقات خيالي ورسالة تلخص النتائج. التقييم: الدقة ٥٠ بالمئة، الوضوح ٣٠ بالمئة، وحماية البيانات ٢٠ بالمئة.')])
]


def render(title, paragraphs, target, language='en'):
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for border in list(doc.styles.element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr")):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.8)
    for name, size in [('Normal',11),('Title',20),('Heading 1',14)]:
        style=doc.styles[name]; style.font.name='Arial'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7)
    doc.add_paragraph(title,'Title')
    doc.add_paragraph(NOTICE)
    html=['<h1>'+escape(title)+'</h1><p>'+escape(NOTICE)+'</p>']
    page_groups=[]
    for heading, text in paragraphs:
        if heading:
            new_page = heading in ('Course content', 'Capstone specification')
            if new_page:
                page_groups.append(''.join(html)); html=[]
            hp=doc.add_paragraph(heading,'Heading 1')
            hp.paragraph_format.page_break_before=new_page
            html.append('<h2>'+escape(heading)+'</h2>')
        # Keep each source field in a separate paragraph for ingestion.
        for line in text.splitlines():
            p=doc.add_paragraph(line)
            p.paragraph_format.keep_together=True
            rtl=any('\u0600'<=c<='\u06ff' for c in line)
            if rtl:
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                flag=OxmlElement('w:bidi'); flag.set(qn('w:val'),'1'); p._p.get_or_add_pPr().append(flag)
            html.append('<p dir="'+('rtl' if rtl else 'ltr')+'">'+escape(line)+'</p>')
    doc.core_properties.author='REX AI demo fixtures'
    sec.footer.paragraphs[0].text='Fictional demonstration material | REX AI'
    doc.save(target.with_suffix('.docx'))
    page_groups.append(''.join(html))
    writer=fitz.DocumentWriter(str(target.with_suffix('.pdf')))
    rect=fitz.Rect(0,0,612,792)
    try:
        for group in page_groups:
            story=fitz.Story(html=group,user_css='body{font-family:sans-serif;font-size:11pt;line-height:1.35;} h1{font-size:20pt;}h2{font-size:14pt;}p{margin:0 0 7pt;}')
            more, filled=story.place(rect+(52,48,-52,-48))
            if more:
                raise RuntimeError('Fixture page overflow: revise content or layout')
            device=writer.begin_page(rect)
            story.draw(device)
            writer.end_page()
    finally: writer.close()



def build():
    OUT.mkdir(parents=True,exist_ok=True)
    manifest=[]
    labels={'course_id':'Course ID','title':'Title','description':'Description','level':'Level','training_format':'Training format','price_egp':'Price EGP','duration_hours':'Duration hours','department':'Department','instructor_name':'Instructor name','instructor_bio':'Instructor bio','schedule':'Schedule','offers':'Offers','target_audience':'Target audience','career_opportunities':'Career opportunities','admission_requirements':'Admission requirements','prerequisites':'Prerequisites','learning_outcomes':'Learning outcomes','faqs':'FAQs'}
    for c in COURSES:
        fields={k:v for k,v in c.items() if k!='modules'}
        lines=[]
        for key,label in labels.items():
            value=fields[key]
            if key=='faqs': value=json.dumps(value,ensure_ascii=False)
            elif isinstance(value,list): value='; '.join(value)
            elif key=='price_egp': value=str(value)+' EGP'
            lines.append(f'{label}: {value}')
        text='\n'.join(lines)
        paragraphs=[('',text),('Course content','This material is an authored fixture for testing ingestion and document workflows. It is not an accredited curriculum.')]+c['modules']
        target=OUT/'courses'/c['course_id'].lower()
        target.parent.mkdir(parents=True,exist_ok=True)
        target.with_suffix('.md').write_text('# '+c['title']+'\n\n'+NOTICE+'\n\n'+text+'\n\n'+'\n\n'.join('## '+h+'\n'+t for h,t in paragraphs[1:]),encoding='utf-8')
        render(c['title'],paragraphs,target,'ar' if c['course_id']=='DEMO-AR-101' else 'en')
        manifest.append({'file':'courses/'+target.name+'.md','expected':fields})
    (OUT/'course_manifest.json').write_text(json.dumps({'data_origin':'synthetic_demo','notice':NOTICE,'courses':manifest},ensure_ascii=False,indent=2),encoding='utf-8')
    refs=[
      ('Participant guide','For the Python Foundations Demo learner. Before class install Python and create a practice folder. After each module complete the course exercise, record one difficulty and ask for feedback. Final task: submit a CSV expense tracker with a README and three validation tests.'),
      ('Facilitator guide','Authored demonstration reference. Session structure: explain for 20 minutes, demonstrate for 20, practice for 90, review for 30 and reflect for 20. Check understanding using a small numeric example before introducing CSV input. Support learners with a worked example, then remove the scaffolding.'),
      ('Presentation outline','Slide 1: learning goal and fictional context. Slide 2: data types and conversion. Slide 3: rejecting negative values. Slide 4: totaling expenses with a loop. Slide 5: reusable functions. Slide 6: reading CSV safely. Slide 7: expense tracker demonstration. Slide 8: reflection and next steps. This is an outline, not a generated slide deck.'),
      ('Question bank','Q1: What is sum([120,80,50])? Answer: 250. Q2: Why validate float input? Answer: invalid input must produce a clear error rather than crash. Q3: When should a function be used? Answer: to isolate and reuse a coherent operation. Practical question: implement a total function that rejects negative expenses.'),
      ('Assessment rubric','Correctness 50 points: totals and CSV round trip work. Validation 30 points: missing, nonnumeric and negative values handled clearly. Clarity 20 points: meaningful names, a short README and repeatable test steps. Passing demonstration threshold: 60 out of 100. This threshold is an authored fixture, not academy policy.'),
      ('Capstone specification','Build a local expense tracker with add, list and total commands. Use only fictional transactions. Persist a date, category and amount in CSV. Deliver source code, test inputs, expected totals and a README. Exclude authentication, payments and external personal data from this learning exercise.'),
      ('Course description','Python Foundations Demo is a fictional 24-hour beginner online course at a demo price of 2400 EGP. It teaches variables, conditions, loops, functions and CSV handling. No programming experience is required. No job, certificate or accreditation is promised.'),
      ('Curriculum example','Authored reference fixture only; the Curriculum Agent is not implemented by this file. Allocate 6 hours each to getting started, control flow, functions and files, and the expense tracker project. Map each module exercise to the corresponding learning outcome and collect the final project for review.'),
      ('Training proposal','Fictional client: Cedar Demo Company. Need: introduce ten office staff to Python automation. Proposed scope: the 24-hour Python Foundations Demo course, online delivery, exercises and a final tracker. Demo tuition for ten learners before any eligible demo discount: 24000 EGP. Venue, dates and acceptance require human confirmation. This is not a commercial quote.'),
      ('Learning outcomes and assessment plan','Outcome 1 write conditions and loops: inspect a totaling exercise. Outcome 2 read and validate CSV: test missing and malformed rows. Outcome 3 build a tracker: review the capstone and rubric. Collect formative feedback every session and use the final project for a summative assessment. No external accreditation claim is made.')]
    render('Teaching document reference pack',refs,OUT/'reference_materials'/'teaching_reference_pack')
    (OUT/'reference_materials'/'teaching_reference_pack.md').write_text('# Teaching document reference pack\n\n'+NOTICE+'\n\n'+'\n\n'.join('## '+h+'\n'+t for h,t in refs),encoding='utf-8')
    neg=OUT/'negative_cases';neg.mkdir(exist_ok=True)
    (neg/'reject_missing_information.md').write_text('# Incomplete course example\n\n'+NOTICE+'\n\nTitle: Incomplete Demo\nDescription: This fixture intentionally omits a course ID, price, level and prerequisites. Reject it with a clear reason; never invent missing details.',encoding='utf-8')
    (neg/'multiple_courses.md').write_text('## First fictional course\n'+NOTICE+'\nCourse ID: DEMO-MULTI-1\nTitle: First Demo\nDescription: First independent example.\nPrice EGP: 1000 EGP\n\n## Second fictional course\nCourse ID: DEMO-MULTI-2\nTitle: Second Demo\nDescription: Second independent example.\nPrice EGP: 2000 EGP\n',encoding='utf-8')
    (OUT/'learner_profiles.json').write_text(json.dumps({'notice':NOTICE,'examples':[{'goal':'Learn Python for office automation','level':'Beginner','budget':'3000 EGP','format':'Online','expected':'DEMO-PY-101 is a candidate; prerequisites still require review.'},{'goal':'Cloud operations','level':'Beginner','budget':'1000 EGP','format':'Online','expected':'No suitable cloud match. Do not recommend the advanced onsite course.'},{'goal':'Data analysis','level':'Intermediate','budget':'6000 EGP','format':'Hybrid','expected':'DEMO-DA-201 is a candidate.'}]},indent=2),encoding='utf-8')
    # Relative dates are refreshed by the staging endpoint when the TA loads the pack.
    rows=[]
    for i,(country,title) in enumerate([('Egypt','Python staff training'),('Saudi Arabia','Data analysis capacity building'),('Germany','Software training'),('Egypt','Expired cloud training')],1):
        rows.append({'title':title+' REX_SYNTHETIC_DEMO','issuer':'Cedar Demo Organisation '+str(i),'country':country,'publication_date':str(date.today()-timedelta(days=2)),'deadline':str(date.today()+timedelta(days=-1 if i==4 else 60)),'category':'IT training','description':NOTICE+' Training and professional development for fictional employees.','source_url':f'https://example.invalid/rex-demo/tender/{i}','submission_url':f'https://example.invalid/rex-demo/apply/{i}','eligibility':'Fictional training provider only','data_mode':'synthetic_demo'})
    rows.append(dict(rows[0]))
    (OUT/'tenders.json').write_text(json.dumps({'notice':NOTICE,'items':rows},indent=2),encoding='utf-8')
    labs=OUT/'course_content'; labs.mkdir(exist_ok=True)
    (labs/'expenses.csv').write_text('date,category,amount,data_origin\n2026-10-01,books,120,REX_SYNTHETIC_DEMO\n2026-10-02,transport,80,REX_SYNTHETIC_DEMO\n2026-10-03,stationery,50,REX_SYNTHETIC_DEMO\n',encoding='utf-8')
    (labs/'expenses_invalid.csv').write_text('date,category,amount,data_origin\n2026-10-01,books,120,REX_SYNTHETIC_DEMO\n2026-10-02,transport,-15,REX_SYNTHETIC_DEMO\n2026-10-03,stationery,apple,REX_SYNTHETIC_DEMO\n2026-10-04,books,,REX_SYNTHETIC_DEMO\n',encoding='utf-8')
    (labs/'sales.csv').write_text('order_id,category,quantity,unit_price,data_origin\nD001,laptops,2,1500,REX_SYNTHETIC_DEMO\nD002,cables,3,100,REX_SYNTHETIC_DEMO\nD003,mice,1,250,REX_SYNTHETIC_DEMO\nD001,laptops,2,1500,REX_SYNTHETIC_DEMO\nD004,cables,,100,REX_SYNTHETIC_DEMO\n',encoding='utf-8')
    (labs/'README.md').write_text('# Fictional course exercise inputs\n\n'+NOTICE+'\n\nexpenses.csv: expected total 250 EGP across three rows. expenses_invalid.csv: only the first amount is valid; reject or explicitly report the negative, nonnumeric and missing values. sales.csv: remove the exact duplicate D001 row, flag D004 for missing quantity, and calculate revenue for the three complete unique rows as 3550 EGP. Do not silently fill the missing quantity or claim the rows describe real sales.\n',encoding='utf-8')
    print('Created four course packs, a ten-type authored reference pack, tender and learner fixtures, and negative cases in',OUT)

if __name__=='__main__': build()
