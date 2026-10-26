"""Deterministic DOCX/PDF rendering of persisted, source-cited document content."""
from html import escape
import re
from pathlib import Path

ARABIC = re.compile(r"[\u0600-\u06ff]")


def _blocks(record):
    ar = record["language"] == "ar"
    yield "title", record["content"]["title"]
    if record.get("data_origin") == "synthetic_demo":
        yield "meta", "REX_SYNTHETIC_DEMO | FICTIONAL DEMONSTRATION | NOT A REAL ACADEMY OFFERING"
    label = (("معتمد بعد المراجعة البشرية" if ar else "Approved after human review") if record["status"] == "approved"
             else ("مسودة للمراجعة البشرية" if ar else "Draft for human review"))
    yield "meta", label + f' | REX AI | v{record["version"]}'
    if record.get("approved_by"):
        yield "meta", f'{record["approved_by"]} | {record["approved_at"]}'
    yield "body", ("الجمهور المستهدف" if ar else "Intended audience") + ": " + record["audience"]
    for section in record["content"]["sections"]:
        yield "heading", section["heading"]
        yield "body", section["text"]
        yield "meta", " ".join(f"[{sid}]" for sid in section["source_ids"])
    yield "heading", "بيانات الدورات المعتمدة" if ar else "Verified course details"
    labels = {"title": ("Course", "الدورة"), "description": ("Description", "الوصف"),
              "level": ("Level", "المستوى"), "training_format": ("Delivery", "نظام التدريب"),
              "price_egp": ("Price EGP", "السعر بالجنيه المصري"), "duration_hours": ("Duration hours", "عدد الساعات"),
              "department": ("Department", "القسم"), "prerequisites": ("Prerequisites", "المتطلبات السابقة"),
              "learning_outcomes": ("Learning outcomes", "مخرجات التعلم")}
    for source in record["sources"]:
        yield "heading", f'[{source["source_id"]}] {source["course"]["title"]}'
        for field, pair in labels.items():
            value = source["course"].get(field)
            if value is None or value == []:
                value = "غير محدد في البيانات المعتمدة" if ar else "Not specified in approved data"
            elif isinstance(value, list):
                value = "; ".join(value)
            yield "body", pair[1 if ar else 0] + ": " + str(value)
        yield "meta", f'{source["source_id"]} | {source["course"]["course_id"]} | {source["source_file"]}'
        yield "meta", "SHA256 " + source["course"]["source_sha256"]
    yield "meta", "Artifact " + record["artifact_id"]


def render_document(record, directory: Path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import pymupdf as fitz

    directory.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for border in list(doc.styles.element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr")):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else 21 if name == "Title" else 14)
        style.paragraph_format.space_after = Pt(7)
    doc.core_properties.author = "REX AI"
    doc.core_properties.title = record["content"]["title"]
    html = []
    for kind, text in _blocks(record):
        style = {"title": "Title", "heading": "Heading 1", "body": "Normal", "meta": "Normal"}[kind]
        paragraph = doc.add_paragraph(style=style)
        rtl = bool(ARABIC.search(text))
        if rtl:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            flag = OxmlElement("w:bidi")
            flag.set(qn("w:val"), "1")
            paragraph._p.get_or_add_pPr().append(flag)
        run = paragraph.add_run(text)
        if kind == "meta":
            run.font.size = Pt(8)
        if rtl:
            run.font.complex_script = True
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:cs"), "Arial")
        tag = {"title": "h1", "heading": "h2", "body": "p", "meta": "p"}[kind]
        html.append(f'<{tag} class="{kind}" dir="{"rtl" if rtl else "ltr"}">{escape(text).replace(chr(10), "<br/>")}</{tag}>')
    footer = sec.footer.paragraphs[0]
    footer.add_run("REX AI | " + record["artifact_id"])
    footer.runs[0].font.size = Pt(8)
    doc.save(directory / "document.docx")
    # MuPDF Story uses HarfBuzz shaping and bundled fallback fonts for Arabic.
    css = """body {font-family: sans-serif; font-size: 10pt; line-height:1.4; color:#111;}
    h1 {font-size:21pt; margin-bottom:10pt;} h2 {font-size:14pt; margin-top:17pt; page-break-after:avoid;}
    p {margin-top:0; margin-bottom:8pt;} .meta {font-size:8pt; color:#444; overflow-wrap:anywhere;}
    """
    story = fitz.Story(html="".join(html), user_css=css)
    writer = fitz.DocumentWriter(str(directory / "document.pdf"))
    mediabox = fitz.paper_rect("a4")
    def rectfn(rect_num, filled):
        return mediabox, mediabox + (48, 48, -48, -48), None
    try:
        story.write(writer, rectfn)
    finally:
        writer.close()
    with fitz.open(directory / "document.pdf") as pdf:
        if not len(pdf) or any(not page.get_text().strip() for page in pdf):
            raise RuntimeError("PDF rendering produced an empty page")
