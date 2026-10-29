"""Shared Andalusia Academy design language for PDF and editable Word exports."""
from html import escape
import io
import re
from copy import deepcopy

STYLE_VERSION = "andalusia-official-portfolio-2"
BURGUNDY = "6F1D1D"
COPPER = "B07963"
IVORY = "FAF6F3"
INK = "212D30"
MUTED = "666666"
ARABIC = re.compile(r"[\u0600-\u06ff\ufb50-\ufeff]")


def _html(text):
    return escape(str(text)).replace("\n", "<br/>")


def _direction(text):
    return "rtl" if ARABIC.search(str(text)) else "ltr"


def logo_png():
    from pathlib import Path
    import pymupdf
    path = Path(__file__).resolve().parents[3] / "assets/branding/andalusia-logo.svg"
    with pymupdf.open(stream=path.read_bytes(), filetype="svg") as image:
        return image[0].get_pixmap(matrix=pymupdf.Matrix(3, 3), alpha=True).tobytes("png")


def status_label(status, ar=False):
    if status == "approved":
        return "معتمد بعد المراجعة البشرية" if ar else "Approved after human review"
    return "مسودة للمراجعة البشرية" if ar else "Draft for human review"


def render_pdf(blocks, *, title, kind, status, version=1, demo=False, ar=False, cover_overview="", summary=None):
    """Shape Unicode with MuPDF Story, then add consistent page furniture."""
    import pymupdf
    if not cover_overview:
        cover_overview = next((b.get("text", "") for b in blocks if b["kind"] == "paragraph"), "")
    overview = cover_overview[:420].rsplit(" ", 1)[0] + "…" if len(cover_overview) > 420 else cover_overview
    summary = summary or []
    metric_cells = ''.join(f'<td style="width:{420/max(1,len(summary))}pt"><p class="metric-value">{_html(value)}</p><p class="metric-label">{_html(label)}</p></td>' for label,value in summary)
    cover = '<div class="cover"><p class="cover-space"></p>'
    cover += f'<p class="cover-kicker">{_html(kind.upper())}</p><h1 class="cover-title" dir="{_direction(title)}">{_html(title)}</h1>'
    cover += f'<p class="cover-status">{_html(status_label(status,ar))} · v{version}</p>'
    if overview:
        cover += f'<div class="cover-overview" dir="{_direction(overview)}"><p class="cover-kicker">{"نبذة عن المستند" if ar else "DOCUMENT OVERVIEW"}</p><p>{_html(overview)}</p></div>'
    if summary:
        cover += f'<table class="cover-stats"><tr>{metric_cells}</tr></table>'
    if demo:
        cover += '<p class="demo">REX_SYNTHETIC_DEMO · FICTIONAL DEMONSTRATION</p>'
    cover += '</div>'
    pieces = [f'<p class="eyebrow">{_html(kind)}</p>']
    segments = [[cover]]
    for block in blocks:
        text = block.get("text", "")
        direction = _direction(text)
        kind_ = block["kind"]
        if kind_ == "appendix":
            segments.append(pieces)
            pieces = []
            pieces.append(f'<h2 class="appendix" dir="{direction}">{_html(text)}</h2>')
        elif kind_ in ("section", "subheading"):
            tag = "h2" if kind_ == "section" else "h3"
            pieces.append(f'<{tag} dir="{direction}">{_html(text)}</{tag}>')
        elif kind_ == "facts":
            rows = ''.join(f'<tr><td class="label" style="width:145pt" dir="{_direction(k)}">{_html(k)}</td><td style="width:315pt" dir="{_direction(v)}">{_html(v)}</td></tr>' for k, v in block["rows"])
            if ar:
                rows = ''.join(f'<tr><td style="width:315pt" dir="{_direction(v)}">{_html(v)}</td><td class="label" style="width:145pt" dir="rtl">{_html(k)}</td></tr>' for k, v in block["rows"])
                pieces.append('<table class="facts"><thead><tr><th style="width:315pt">المعلومات</th><th style="width:145pt">البيان</th></tr></thead><tbody>' + rows + '</tbody></table>')
            else:
                pieces.append('<table class="facts"><thead><tr><th style="width:145pt">Detail</th><th style="width:315pt">Information</th></tr></thead><tbody>' + rows + '</tbody></table>')
        elif kind_ == "bullets":
            if ar:
                pieces.extend(f'<p class="arabic-bullet" dir="rtl">• {_html(v)}</p>' for v in block["items"])
            else:
                items = ''.join(f'<li dir="{_direction(v)}">{_html(v)}</li>' for v in block["items"])
                pieces.append(f'<ul>{items}</ul>')
        else:
            cls = "reference" if kind_ == "reference" else "callout" if kind_ == "callout" else "body"
            for paragraph in str(text).split("\n\n"):
                pieces.append(f'<p class="{cls}" dir="{_direction(paragraph)}">{_html(paragraph)}</p>')
    css = f'''
    body{{margin:0;font-family:sans-serif;font-size:10.5pt;line-height:1.52;color:#{INK};}}
    h1{{font-family:sans-serif;font-weight:bold;font-size:28pt;line-height:1.18;color:#{BURGUNDY};margin:8pt 0 16pt;}}
    h2{{font-size:15pt;line-height:1.25;color:white;background:#{BURGUNDY};padding:9pt 11pt;margin:21pt 0 9pt;page-break-after:avoid;}}
    h3{{font-size:11.5pt;line-height:1.3;color:#{BURGUNDY};margin:14pt 0 7pt;page-break-after:avoid;}}
    p{{margin:0 0 10pt;}} li{{margin:0 0 5pt;}} ul{{margin:4pt 0 12pt;padding-left:19pt;}}
    .eyebrow{{font-size:9pt;color:#{MUTED};margin:0 0 6pt;}}
    .review{{background:#{IVORY};color:#{BURGUNDY};font-size:9pt;padding:10pt 12pt;margin-bottom:15pt;}}
    .demo{{font-size:8pt;color:#795a26;background:#fbf5e6;padding:8pt;}}
    .callout{{background:#{IVORY};padding:11pt 13pt;font-size:10pt;margin:12pt 0;}}
    .reference{{font-size:8pt;line-height:1.4;color:#{MUTED};overflow-wrap:anywhere;margin-bottom:7pt;}}
    .appendix{{margin-top:0;padding:0;background:#ffffff;color:#{BURGUNDY};}}
    table{{width:100%;border-collapse:collapse;margin:10pt 0 15pt;font-size:9pt;}}
    th{{background:#{BURGUNDY};color:white;font-size:9pt;text-align:left;padding:8pt;}}
    td{{border:0.5pt solid #d9dfd9;padding:8pt;vertical-align:top;}}
    .label{{background:#{IVORY};font-weight:bold;}}
    .arabic-bullet{{margin-right:10pt;margin-bottom:6pt;}}
    .cover-space{{height:100pt;margin:0;}}
    .cover{{text-align:center;}}
    .cover-kicker{{font-size:8pt;letter-spacing:1.4pt;color:#{COPPER};margin-bottom:14pt;}}
    .cover-title{{font-size:30pt;line-height:1.17;text-align:center;color:#{BURGUNDY};margin:12pt 0 20pt;}}
    .cover-status{{font-size:9pt;color:#{MUTED};margin-bottom:24pt;}}
    .cover-overview{{background:#{IVORY};border:.6pt solid #ddc3b4;padding:18pt 20pt;margin:12pt 32pt 22pt;font-size:10pt;line-height:1.6;}}
    .cover-stats{{margin:20pt 20pt;}} .cover-stats td{{text-align:center;border:.6pt solid #ddc3b4;padding:13pt 8pt;}}
    .metric-value{{font-size:19pt;font-weight:bold;color:#{BURGUNDY};margin-bottom:5pt;}}
    .metric-label{{font-size:8pt;color:#{MUTED};margin:0;}}
    '''
    segments.append(pieces)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page_rect = pymupdf.paper_rect("a4")
    device = None
    y = 88
    pages = 0
    def new_page():
        nonlocal device, y, pages
        if device is not None:
            writer.end_page()
            device = None
        pages += 1
        if pages > 500:
            raise RuntimeError("Document exceeds the page limit")
        device = writer.begin_page(page_rect)
        y = 88
    try:
        for segment in segments:
            new_page()
            groups, pending = [], []
            for piece in segment:
                if groups and piece.startswith('<p class="reference"') and ("Sources:" in piece or "المراجع:" in piece):
                    groups[-1] += piece
                    continue
                pending.append(piece)
                if not piece.startswith(("<h2", "<h3", "<h1", '<p class="eyebrow"')):
                    groups.append(''.join(pending)); pending = []
            if pending:
                groups.append(''.join(pending))
            for group in groups:
                if y > page_rect.height - 91:
                    new_page()
                # Layout each block independently. A shaded table that fits a
                # page moves intact, avoiding repeated background painting when
                # Story flows an earlier table through subsequent text pages.
                story = pymupdf.Story(html='<html><body>' + group + '</body></html>', user_css=css)
                more, filled = story.place(pymupdf.Rect(48, y, page_rect.width - 48, page_rect.height - 61))
                if more and y > 88:
                    story.reset()
                    new_page()
                    more, filled = story.place(pymupdf.Rect(48, y, page_rect.width - 48, page_rect.height - 61))
                while True:
                    if more and filled[3] <= y:
                        raise RuntimeError("Document block cannot fit the page")
                    story.draw(device)
                    y = filled[3] + 2
                    if not more:
                        break
                    new_page()
                    more, filled = story.place(pymupdf.Rect(48, y, page_rect.width - 48, page_rect.height - 61))
        if device is not None:
            writer.end_page()
            device = None
    finally:
        if device is not None:
            writer.end_page()
        writer.close()
    green = tuple(int(BURGUNDY[i:i+2], 16) / 255 for i in (0, 2, 4))
    gold = tuple(int(COPPER[i:i+2], 16) / 255 for i in (0, 2, 4))
    with pymupdf.open(stream=buffer.getvalue(), filetype="pdf") as content_pdf, pymupdf.open() as pdf:
        mark = logo_png()
        for i in range(len(content_pdf)):
            # Isolate the Story's graphics state in a form XObject so RTL text
            # clipping cannot hide the independently drawn running header.
            page = pdf.new_page(width=page_rect.width, height=page_rect.height)
            page.draw_rect(pymupdf.Rect(0, 0, page.rect.width, 5), color=green, fill=green)
            if i == 0:
                page.insert_image(pymupdf.Rect(200, 121, 395, 168), stream=mark)
            else:
                page.insert_image(pymupdf.Rect(48, 29, 178, 60), stream=mark)
                page.insert_text((335, 47), "ANDALUSIA ACADEMY / REX AI", fontsize=7, color=gold)
            page.draw_line((48, page.rect.height - 45), (page.rect.width - 48, page.rect.height - 45), color=gold, width=.6)
            page.insert_text((48, page.rect.height - 29), f"{'DEMO  |  ' if demo else ''}{'APPROVED' if status == 'approved' else 'DRAFT'}  |  v{version}", fontsize=7, color=green)
            page.insert_text((page.rect.width - 105, page.rect.height - 29), f"{i+1:02d} / {len(content_pdf):02d}", fontsize=8, color=green)
            page.show_pdf_page(page.rect, content_pdf, i, overlay=False)
        pdf.set_metadata({"title": title, "author": "Andalusia Academy", "subject": kind, "keywords": STYLE_VERSION})
        return pdf.tobytes(garbage=3, deflate=True)


def render_word(blocks, *, title, kind, status, version=1, demo=False, ar=False, cover_overview="", summary=None):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin, sec.bottom_margin = Inches(1.1), Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.72)
    sec.header_distance = Inches(.32)
    sec.footer_distance = Inches(.35)
    for name, size, color in (("Normal", 11, INK), ("Title", 29, BURGUNDY), ("Heading 1", 16, BURGUNDY), ("Heading 2", 12, BURGUNDY), ("Caption", 9, MUTED)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
            fonts.attrib.pop(qn("w:" + attribute), None)
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(9)
        style.paragraph_format.widow_control = True
        if name == "Caption":
            style.font.bold = False
            style.font.italic = False
        if name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(17 if name == "Heading 1" else 11)
            style.paragraph_format.keep_with_next = True
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    def paragraph(text, style="Normal", *, shade=None):
        p = doc.add_paragraph(style=style)
        rtl = bool(ARABIC.search(str(text)))
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
        if rtl:
            flag = OxmlElement("w:bidi"); flag.set(qn("w:val"), "1")
            p._p.get_or_add_pPr().append(flag)
        run = p.add_run(str(text))
        if rtl:
            props = run._element.get_or_add_rPr()
            flag = OxmlElement("w:rtl"); flag.set(qn("w:val"), "1"); props.append(flag)
            props.get_or_add_rFonts().set(qn("w:cs"), "Arial")
            font_size = doc.styles[style].font.size or doc.styles["Normal"].font.size
            size = OxmlElement("w:szCs"); size.set(qn("w:val"), str(int(font_size.pt * 2))); props.append(size)
        if shade:
            fill = OxmlElement("w:shd"); fill.set(qn("w:fill"), shade); p._p.get_or_add_pPr().append(fill)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(13)
        return p
    header = sec.header.paragraphs[0]
    header.add_run().add_picture(io.BytesIO(logo_png()), width=Inches(1.75))
    r = header.add_run("       ANDALUSIA ACADEMY / REX AI")
    r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(COPPER)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run(f"Andalusia Academy  |  {'DEMO  |  ' if demo else ''}{'APPROVED' if status == 'approved' else 'DRAFT'}  |  v{version}     ")
    r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    # The cover follows the supplied portfolio reference. Ordinary running
    # headers resume after the cover; the official logo is never redrawn.
    sec.different_first_page_header_footer = True
    sec.first_page_footer._element.append(deepcopy(footer._p))
    mark = paragraph("")
    mark.paragraph_format.space_before = Pt(66)
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark.add_run().add_picture(io.BytesIO(logo_png()), width=Inches(2.15))
    kicker = paragraph(kind.upper(), "Caption")
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(25)
    for run in kicker.runs:
        run.font.color.rgb = RGBColor.from_string(COPPER)
    cover_title = paragraph(title, "Title")
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_title.paragraph_format.space_after = Pt(20)
    for run in cover_title.runs:
        run.font.bold = True
        run.font.size = Pt(24 if len(title) > 130 else 29)
    p = paragraph(status_label(status, ar) + f"  |  v{version}", "Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not cover_overview:
        cover_overview = next((b.get("text", "") for b in blocks if b["kind"] == "paragraph"), "")
    overview = cover_overview[:420].rsplit(" ",1)[0] + "…" if len(cover_overview)>420 else cover_overview
    if overview:
        p = paragraph(overview, shade=IVORY)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(.45)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if summary:
        metrics = doc.add_table(rows=1, cols=len(summary))
        metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
        metrics.autofit = False
        for column in metrics.columns:
            column.width = Inches(6.7 / len(summary))
        for cell, (label, value) in zip(metrics.rows[0].cells, summary):
            cell.width = Inches(6.7 / len(summary))
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(str(value) + "\n"); run.font.name = "Arial"; run.font.size = Pt(18)
            run.font.bold = True; run.font.color.rgb = RGBColor.from_string(BURGUNDY)
            run = p.add_run(label); run.font.name = "Arial"; run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideV"):
            el = OxmlElement("w:" + side)
            for key,value in (("val","single"),("sz","4"),("color","DDC3B4")):
                el.set(qn("w:"+key),value)
            borders.append(el)
        metrics._tbl.tblPr.append(borders)
    if demo:
        p = paragraph("REX_SYNTHETIC_DEMO | FICTIONAL DEMONSTRATION | NOT A REAL ACADEMY OFFERING", "Caption", shade=IVORY)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    paragraph(kind, "Caption")
    for block in blocks:
        kind_ = block["kind"]
        if kind_ == "facts":
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            table.columns[0].width, table.columns[1].width = Inches(1.8), Inches(4.9)
            labels = ("البيان", "المعلومات") if ar else ("Detail", "Information")
            for cell, label in zip(table.rows[0].cells, labels):
                cell.text = label
            repeat = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for key, value in block["rows"]:
                row = table.add_row(); row.cells[0].text = str(key); row.cells[1].text = str(value)
            props = table._tbl.tblPr
            if ar:
                flag = OxmlElement("w:bidiVisual"); flag.set(qn("w:val"), "1"); props.append(flag)
            borders = OxmlElement("w:tblBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement("w:" + side)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    el.set(qn("w:" + key), value)
                borders.append(el)
            props.append(borders)
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.width = Inches(1.8 if ci == 0 else 4.9)
                    fill = OxmlElement("w:shd"); fill.set(qn("w:fill"), BURGUNDY if ri == 0 else (IVORY if ri % 2 else "FFFFFF")); cell._tc.get_or_add_tcPr().append(fill)
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "bottom", "left", "right"):
                        el = OxmlElement("w:" + side); el.set(qn("w:w"), "100"); el.set(qn("w:type"), "dxa"); margins.append(el)
                    cell._tc.get_or_add_tcPr().append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.keep_with_next = ri == 0
                        if ARABIC.search(p.text):
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            flag = OxmlElement("w:bidi"); flag.set(qn("w:val"), "1"); p._p.get_or_add_pPr().append(flag)
                        for r in p.runs:
                            r.font.name = "Arial"; r.font.size = Pt(10)
                            r.font.bold = ri == 0 or ci == 0
                            r.font.color.rgb = RGBColor.from_string("FFFFFF" if ri == 0 else INK)
            paragraph("").paragraph_format.space_after = Pt(2)
        elif kind_ == "bullets":
            for item in block["items"]:
                paragraph(item, "List Bullet")
        else:
            style = {"section": "Heading 1", "appendix": "Heading 1", "subheading": "Heading 2", "reference": "Caption"}.get(kind_, "Normal")
            for i, text in enumerate(str(block.get("text", "")).split("\n\n")):
                p = paragraph(text, style, shade=BURGUNDY if kind_ == "section" else IVORY if kind_ == "callout" else None)
                if kind_ == "section":
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                if kind_ == "appendix" and i == 0:
                    p.paragraph_format.page_break_before = True
    doc.core_properties.author = "Andalusia Academy"
    doc.core_properties.title = title
    doc.core_properties.subject = kind
    doc.core_properties.keywords = STYLE_VERSION
    return doc
