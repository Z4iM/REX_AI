"""Offline tests for real course-file staging. No Gemini key required."""
from io import BytesIO

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.app.catalog.upload_router import router


def make_docx_bytes(text="Python and Data Analytics course details"):
    doc = Document()
    doc.add_heading("Course Information", level=1)
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("REX_CATALOG_DATA_DIR", str(tmp_path / "catalog"))
    monkeypatch.setenv("REX_CATALOG_ADMIN_KEY", "test-admin-secret")
    app = FastAPI()
    app.include_router(router)
    return TestClient(app)


H = {"X-API-Key": "test-admin-secret"}


def test_api_is_staff_protected(client):
    assert client.get("/api/v1/catalog/uploads").status_code == 401
    assert client.get("/api/v1/catalog/uploads", headers={"X-API-Key": "wrong"}).status_code == 401


def test_upload_is_saved_but_not_approved_or_indexed(client):
    data = make_docx_bytes()
    resp = client.post("/api/v1/catalog/uploads", headers=H,
                       files={"file": ("course.docx", data,
                                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert resp.status_code == 201, resp.text
    result = resp.json()
    assert result["review_status"] == "pending_review"
    assert result["approval_status"] == "unapproved"
    assert result["chunk_count"] >= 1
    detail = client.get(f'/api/v1/catalog/uploads/{result["upload_id"]}', headers=H).json()
    assert detail["indexing_allowed"] is False
    assert len(detail["chunks"]) == result["chunk_count"]
    assert client.get("/api/v1/catalog/uploads", headers=H).json()["count"] == 1


def test_same_file_is_idempotent(client):
    data = make_docx_bytes()
    def send():
        return client.post("/api/v1/catalog/uploads", headers=H,
                           files={"file": ("course.docx", data)}).json()
    first = send()
    second = send()
    assert first["upload_id"] == second["upload_id"]
    assert first["already_uploaded"] is False
    assert second["already_uploaded"] is True
    assert client.get("/api/v1/catalog/uploads", headers=H).json()["count"] == 1


def test_invalid_extension_rejected(client):
    resp = client.post("/api/v1/catalog/uploads", headers=H,
                       files={"file": ("bad.exe", b"not allowed")})
    assert resp.status_code == 415


def test_bad_pdf_rejected(client):
    resp = client.post("/api/v1/catalog/uploads", headers=H,
                       files={"file": ("fake.pdf", b"not pdf")})
    assert resp.status_code == 422
