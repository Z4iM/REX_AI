
import pytest
import pymupdf

from docx import Document

from backend.app.ingestion.document_loader import (
    load_document,
)


def test_unified_pdf_loader(tmp_path):
    pdf_path = tmp_path / "course.pdf"

    with pymupdf.open() as pdf:
        page = pdf.new_page()

        page.insert_text(
            (72, 72),
            "Python Training Program"
        )

        pdf.save(pdf_path)

    result = load_document(pdf_path)

    assert result["file_type"] == "pdf"
    assert result["source_file"] == "course.pdf"
    assert result["approval_status"] == "unapproved"
    assert "Python Training Program" in (
        result["sections"][0]["content"]
    )


def test_unified_docx_loader(tmp_path):
    docx_path = tmp_path / "course.docx"

    document = Document()
    document.add_heading("Data Science", level=1)
    document.add_paragraph("Machine learning course")
    document.save(docx_path)

    result = load_document(docx_path)

    assert result["file_type"] == "docx"
    assert result["source_file"] == "course.docx"
    assert result["approval_status"] == "unapproved"
    assert result["sections"][0]["title"] == "Data Science"


def test_unified_markdown_loader(tmp_path):
    md_path = tmp_path / "course.md"

    md_path.write_text(
        "# Course Catalog\n"
        "Dataset Status: SYNTHETIC_DEMO\n\n"
        "## Python Course\n"
        "Course ID: CRS-001\n"
        "Price: 3200 EGP\n",
        encoding="utf-8",
    )

    result = load_document(md_path)

    assert result["file_type"] == "md"
    assert result["source_file"] == "course.md"
    assert result["sections"][0]["title"] == "Python Course"
    assert result["sections"][0]["metadata"]["course_id"] == (
        "CRS-001"
    )

    # Preserve the legacy Markdown loader's dataset status.
    assert result["approval_status"] == "SYNTHETIC_DEMO"


def test_unified_markdown_mode_pdf(tmp_path):
    pdf_path = tmp_path / "markdown_mode.pdf"

    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "AI Training")
        pdf.save(pdf_path)

    result = load_document(
        pdf_path,
        pdf_extraction_mode="markdown",
    )

    assert result["file_type"] == "pdf"
    assert result["extraction_mode"] == "markdown"
    assert "AI Training" in result["sections"][0]["content"]


def test_unsupported_format(tmp_path):
    file_path = tmp_path / "course.exe"
    file_path.write_text("unsupported")

    with pytest.raises(
        ValueError,
        match="Unsupported file format",
    ):
        load_document(file_path)


def test_missing_document(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_document(tmp_path / "missing.pdf")

def test_unified_loader_scanned_pdf_ocr(tmp_path):
    """
    Verify that the unified loader forwards OCR options
    and recognizes text from a scanned PDF.
    """

    pdf_path = tmp_path / "scanned_training.pdf"

    # Create a normal PDF page.
    with pymupdf.open() as source:

        page = source.new_page()

        page.insert_text(
            (72, 120),
            "Data Science Training Program",
            fontsize=24,
        )

        page.insert_text(
            (72, 170),
            "Duration 100 Hours",
            fontsize=20,
        )

        # Rasterize the page into an image.
        pixmap = page.get_pixmap(
            matrix=pymupdf.Matrix(2, 2),
            alpha=False,
        )

        image_data = pixmap.tobytes("png")

    # Create an image-only PDF.
    with pymupdf.open() as scanned:

        page = scanned.new_page()

        page.insert_image(
            page.rect,
            stream=image_data,
        )

        scanned.save(pdf_path)

    # Confirm there is no selectable text.
    with pymupdf.open(pdf_path) as check:
        assert not check[0].get_text().strip()

    # Verify OCR is available through load_document.
    result = load_document(
        pdf_path,
        enable_ocr=True,
        ocr_language="eng",
    )

    assert result["file_type"] == "pdf"
    assert result["ocr_pages"] == [1]
    assert result["pages_without_text"] == []

    section = result["sections"][0]

    assert section["metadata"]["extraction_method"] == "ocr"

    assert "Data Science Training Program" in (
        section["content"]
    )

    assert "100 Hours" in section["content"]


def test_unified_loader_without_ocr_on_scan(tmp_path):
    """
    Confirm OCR remains optional and disabled by default.
    """

    pdf_path = tmp_path / "image_only.pdf"

    with pymupdf.open() as source:

        page = source.new_page()
        page.insert_text(
            (72, 100),
            "Python Course",
            fontsize=24,
        )

        pixmap = page.get_pixmap(
            matrix=pymupdf.Matrix(2, 2),
            alpha=False,
        )

        image_data = pixmap.tobytes("png")

    with pymupdf.open() as scanned:

        page = scanned.new_page()
        page.insert_image(
            page.rect,
            stream=image_data,
        )
        scanned.save(pdf_path)

    with pytest.raises(
        ValueError,
        match="No extractable text",
    ):
        load_document(
            pdf_path,
            enable_ocr=False,
        )
