
from copy import deepcopy

import pytest
import pymupdf

from docx import Document

from backend.app.ingestion.document_review import (
    prepare_document_for_review,
)


def create_docx(tmp_path, name, content):
    path = tmp_path / name

    document = Document()
    document.add_heading("Training Program", level=1)
    document.add_paragraph(content)
    document.save(path)

    return path


def test_docx_is_pending_review(tmp_path):
    path = create_docx(
        tmp_path,
        "course.docx",
        "Python Data Analytics Training",
    )

    package = prepare_document_for_review(path)

    assert package["review_status"] == "pending_review"
    assert package["approval_status"] == "unapproved"
    assert package["indexing_allowed"] is False

    assert package["chunk_count"] >= 1
    assert len(package["source_sha256"]) == 64

    assert all(
        chunk["metadata"]["source_sha256"]
        == package["source_sha256"]
        for chunk in package["chunks"]
    )


def test_incomplete_pdf_is_blocked(tmp_path):
    path = tmp_path / "incomplete.pdf"

    with pymupdf.open() as pdf:
        first_page = pdf.new_page()
        first_page.insert_text(
            (72, 72),
            "Data Science Course",
        )

        # A second page with no readable text.
        pdf.new_page()
        pdf.save(path)

    package = prepare_document_for_review(path)

    assert package["missing_pages"] == [2]

    assert package["review_status"] == (
        "blocked_incomplete"
    )

    assert package["indexing_allowed"] is False
    assert package["extraction_review_required"] is True


def test_file_changes_have_different_hashes(tmp_path):
    first_path = create_docx(
        tmp_path,
        "original.docx",
        "Python Training",
    )

    second_path = create_docx(
        tmp_path,
        "updated.docx",
        "Python Training with Machine Learning",
    )

    first = prepare_document_for_review(first_path)
    second = prepare_document_for_review(second_path)

    assert first["source_sha256"] != (
        second["source_sha256"]
    )


def test_unsupported_document_is_rejected(tmp_path):
    path = tmp_path / "course.exe"
    path.write_text("unsupported")

    with pytest.raises(
        ValueError,
        match="Unsupported file format",
    ):
        prepare_document_for_review(path)
