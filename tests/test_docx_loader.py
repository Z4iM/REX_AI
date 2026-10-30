
import pytest

from docx import Document

from backend.app.ingestion.docx_loader import load_docx


def test_docx_extraction(tmp_path):
    """Check headings, paragraphs, and tables."""

    file_path = tmp_path / "course.docx"

    document = Document()

    document.add_heading(
        "Python for Data Analytics",
        level=1,
    )

    document.add_paragraph(
        "This course is designed for beginners."
    )

    table = document.add_table(rows=2, cols=2)

    table.cell(0, 0).text = "Property"
    table.cell(0, 1).text = "Value"

    table.cell(1, 0).text = "Price"
    table.cell(1, 1).text = "3200 EGP"

    document.add_heading(
        "Prerequisites",
        level=1,
    )

    document.add_paragraph(
        "No previous programming experience required."
    )

    document.save(file_path)

    result = load_docx(file_path)

    assert result["source_file"] == "course.docx"
    assert result["approval_status"] == "unapproved"

    assert len(result["sections"]) == 2

    first_section = result["sections"][0]

    assert first_section["title"] == (
        "Python for Data Analytics"
    )

    assert "designed for beginners" in (
        first_section["content"]
    )

    assert "Price | 3200 EGP" in (
        first_section["content"]
    )

    second_section = result["sections"][1]

    assert second_section["title"] == "Prerequisites"

    assert "No previous programming" in (
        second_section["content"]
    )


def test_empty_docx(tmp_path):
    """Empty Word documents should be rejected."""

    file_path = tmp_path / "empty.docx"

    document = Document()
    document.save(file_path)

    with pytest.raises(
        ValueError,
        match="No extractable text",
    ):
        load_docx(file_path)


def test_invalid_file_type(tmp_path):
    """Reject unsupported file extensions."""

    file_path = tmp_path / "example.txt"

    file_path.write_text("Hello")

    with pytest.raises(
        ValueError,
        match="Only DOCX",
    ):
        load_docx(file_path)


def test_missing_docx(tmp_path):
    """Missing documents should raise an error."""

    with pytest.raises(FileNotFoundError):
        load_docx(
            tmp_path / "missing.docx"
        )


def test_arabic_document(tmp_path):
    """Check that Arabic text is preserved."""

    file_path = tmp_path / "arabic_course.docx"

    document = Document()

    document.add_heading(
        "برنامج تحليل البيانات",
        level=1,
    )

    document.add_paragraph(
        "تعلم تحليل البيانات باستخدام بايثون"
    )

    document.save(file_path)

    result = load_docx(file_path)

    assert "برنامج تحليل البيانات" == (
        result["sections"][0]["title"]
    )

    assert "بايثون" in (
        result["sections"][0]["content"]
    )
