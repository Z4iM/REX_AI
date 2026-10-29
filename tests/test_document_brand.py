from io import BytesIO
import copy
import pytest
import pymupdf
from docx import Document
from backend.app.documents.rendering import render_document


def fixture_record(language):
    ar = language == "ar"
    return {"artifact_id": "00000000-0000-4000-8000-000000000001", "version": 3,
            "status": "pending_review", "language": language, "document_type": "participant_guide",
            "audience": "المتعلمون الجدد" if ar else "New learners", "data_origin": "synthetic_demo",
            "content": {"title": "دليل المشارك في دورة البرمجة" if ar else "Programming Participant Guide",
                        "sections": [{"heading": "التحضير للدورة" if ar else "Preparing for the course",
                         "text": ("هذه عينة خيالية للتحقق من جودة تنسيق المستند.\n\n- مراجعة المتطلبات السابقة\n- إعداد جهاز الحاسوب" if ar else
                                  "This fictional example checks the presentation of learning material.\n\n- Review prerequisites\n- Prepare your computer"), "source_ids": ["S1"]}]},
            "sources": [{"source_id": "S1", "source_file": "fictional_course.md", "course": {
                "course_id": "TEST-101", "title": "دورة برمجة خيالية" if ar else "Fictional programming course",
                "description": "مصدر خيالي للاختبار فقط" if ar else "Fictional test source only",
                "duration_hours": 24, "price_egp": 2400, "level": "Beginner", "training_format": "Online",
                "prerequisites": ["مهارات الحاسوب" if ar else "Computer literacy"],
                "learning_outcomes": ["تطبيق أساسيات البرمجة" if ar else "Apply programming fundamentals"],
                "source_sha256": "a" * 64}}]}


@pytest.mark.parametrize("language", ["en", "ar"])
def test_branded_files_keep_review_and_source_evidence(tmp_path, language):
    record = fixture_record(language)
    original = copy.deepcopy(record)
    render_document(record, tmp_path)
    assert record == original  # Design changes must not rewrite approved facts.
    doc = Document(tmp_path / "document.docx")
    assert doc.core_properties.author == "Andalusia Academy"
    assert "ANDALUSIA ACADEMY" in doc.sections[0].header.paragraphs[0].text
    assert doc.sections[0].header._element.xpath(".//w:drawing")
    facts = next(t for t in doc.tables if len(t.rows) > 1)
    assert facts.rows[1].cells[1].text
    assert any("SHA256 " + "a" * 64 in p.text for p in doc.paragraphs)
    assert doc.styles["Heading 1"].paragraph_format.keep_with_next
    assert "PAGE" in doc.sections[0].footer._element.xml
    with pymupdf.open(tmp_path / "document.pdf") as pdf:
        text = "".join(page.get_text() for page in pdf)
        assert len(pdf) >= 2
        assert "a" * 64 in text and "REX_SYNTHETIC_DEMO" in text
        for i, page in enumerate(pdf):
            text = page.get_text()
            assert page.get_images() and "DRAFT" in text
            assert f"{i+1:02d} / {len(pdf):02d}" in text
            assert all(b[0] >= 0 and b[1] >= 0 and b[2] <= page.rect.width + 1 and b[3] <= page.rect.height + 1
                       for b in page.get_text("blocks") if b[6] == 0)
        # Regression: earlier table fills must not reappear behind appendix text.
        appendix = pdf[-1]
        assert not any(d.get("fill") and min(d["fill"]) < .99 and d["rect"].width > 200 and d["rect"].height > 8
                       for d in appendix.get_drawings())


def test_long_document_repeats_brand_and_approval_status(tmp_path):
    record = fixture_record("en")
    record.update(status="approved", approved_by="Fixture Reviewer", approved_at="2026-10-29T10:00:00Z")
    record["content"]["sections"][0]["text"] = "\n\n".join(
        f"Learning activity {i}: " + "Practise the approved learning objective and record the result. " * 5 for i in range(35))
    render_document(record, tmp_path)
    with pymupdf.open(tmp_path / "document.pdf") as pdf:
        assert len(pdf) >= 4
        assert all("APPROVED" in page.get_text() and "DRAFT" not in page.get_text() for page in pdf)
        assert "Learning activity 34:" in "".join(p.get_text() for p in pdf)


def test_html_in_source_is_printed_as_data(tmp_path):
    record = fixture_record("en")
    record["content"]["sections"][0]["text"] = "<script>untrusted content</script> is source text."
    render_document(record, tmp_path)
    with pymupdf.open(tmp_path / "document.pdf") as pdf:
        assert "<script>untrusted content</script>" in "".join(p.get_text() for p in pdf)
