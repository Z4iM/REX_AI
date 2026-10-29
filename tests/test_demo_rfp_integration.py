"""Sample staging integration using temporary, generated test inputs."""
import io
import json
import zipfile
from pathlib import Path
import pytest
from fastapi import HTTPException
from backend.app import runtime
from backend.app.catalog import catalog_service as store, approved_index
from backend.app.integrations import rfp
from test_catalog_approval_publication import FakeCollection, fake_embed

@pytest.fixture(autouse=True)
def example_pack(tmp_path, monkeypatch):
    """Tests must not depend on an operator's external demonstration files."""
    from docx import Document
    import pymupdf
    folder = tmp_path / 'examples'
    courses = folder / 'courses'
    courses.mkdir(parents=True)
    manifest = {'courses': []}
    for index, code in enumerate(('py', 'da', 'cl', 'ar')):
        name = f'demo-{code}-101'
        text = '\n'.join([
            'REX_SYNTHETIC_DEMO', f'Course ID: DEMO-{code.upper()}-101',
            'Title: Python Foundations Demo' if index == 0 else f'Title: Sample course {index}',
            'Description: Fictional course source for isolated integration testing.',
            'Level: Beginner', 'Training format: Online', 'Price EGP: 2400 EGP',
            'Duration hours: 24', 'Instructor name: Nour Demo Instructor',
            'Schedule: Eight fictional sessions, dates require confirmation.',
            'Prerequisites: Basic computer skills', 'Learning outcomes: Use Python',
            'FAQs: [{"question":"Is programming experience required?","answer":"No, this is a fictional beginner example."}]',
        ])
        (courses / f'{name}.md').write_text(text, encoding='utf-8')
        manifest['courses'].append({'file': f'courses/{name}.md'})
        if index == 0:
            word = Document()
            for line in text.splitlines():
                word.add_paragraph(line)
            word.save(courses / f'{name}.docx')
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                assert page.insert_textbox(pymupdf.Rect(40,40,550,800), text, fontsize=10) >= 0
                pdf.save(courses / f'{name}.pdf')
    (folder / 'course_manifest.json').write_text(json.dumps(manifest), encoding='utf-8')
    rows = []
    for index, country in enumerate(('Egypt', 'Saudi Arabia', 'Germany', 'Egypt')):
        rows.append({'title': ('Expired ' if index == 3 else '') + f'Python staff training {index} REX_SYNTHETIC_DEMO',
                     'issuer': f'Fictional organisation {index}', 'country': country,
                     'category': 'IT training', 'description': 'REX_SYNTHETIC_DEMO training and professional development.',
                     'source_url': f'https://example.invalid/tender/{index}', 'data_mode': 'synthetic_demo'})
    rows.append(dict(rows[0]))
    (folder / 'tenders.json').write_text(json.dumps({'items': rows}), encoding='utf-8')
    references = folder / 'reference_materials'
    references.mkdir()
    word = Document()
    word.add_paragraph('REX_SYNTHETIC_DEMO test-only teaching reference')
    word.save(references / 'teaching_reference_pack.docx')
    monkeypatch.setenv('REX_DEMO_EXAMPLES_DIR', str(folder))
    return folder

@pytest.fixture
def demo(monkeypatch, tmp_path):
    monkeypatch.setenv('REX_APP_MODE','demo')
    monkeypatch.setenv('REX_DEMO_DATA_DIR',str(tmp_path/'demo'))


def test_demo_setup_is_authenticated_and_never_available_in_academy(catalog_client, staff_headers, example_pack):
    assert catalog_client.post('/api/v1/demo/stage-courses',json={'reviewer':'TA'}).status_code==401
    assert catalog_client.post('/api/v1/demo/stage-courses',headers=staff_headers,json={'reviewer':'TA'}).status_code==409
    assert catalog_client.get('/api/v1/demo/examples.zip',headers=staff_headers).status_code==409
    raw=(example_pack/'courses/demo-py-101.md').read_bytes()
    response=catalog_client.post('/api/v1/catalog/uploads',headers=staff_headers,files={'file':('demo.md',raw)})
    assert response.status_code==409
    assert store.list_courses(staff=True)==[]


def test_demo_storage_cannot_use_configured_academy_paths(demo, monkeypatch, tmp_path):
    monkeypatch.setenv('REX_CATALOG_DATA_DIR',str(tmp_path/'real-catalog'))
    monkeypatch.setenv('RFP_DATABASE_URL','sqlite:///real-rfp.db')
    from backend.app.agents.rfp import config
    assert runtime.catalog_dir()==tmp_path/'demo/catalog'
    assert 'real-rfp' not in config.database_url()
    assert runtime.rfp_dir()==tmp_path/'demo/rfp'


def stage(client,headers):
    response=client.post('/api/v1/demo/stage-courses',headers=headers,json={'reviewer':'TA Reviewer'})
    assert response.status_code==200,response.text
    return response.json()['items']


def test_stage_review_publish_rag_reject_and_repeat(demo,catalog_client,staff_headers):
    items=stage(catalog_client,staff_headers)
    assert len(items)==4
    assert catalog_client.get('/api/v1/catalog/courses').json()['count']==0
    primary=catalog_client.get('/api/v1/catalog/uploads/'+items[0]['upload_id'],headers=staff_headers).json()
    candidates=primary['metadata_suggestions']['candidates']
    fields=next(c['fields'] for c in candidates if c['fields'].get('course_id')=='DEMO-PY-101')
    assert fields['instructor_name']=='Nour Demo Instructor'
    assert fields['faqs'][0]['question']=='Is programming experience required?'
    assert fields['price_egp']==2400 and fields['schedule']
    fields.update(reviewer='TA Reviewer',acknowledge_extraction_warnings=True)
    response=catalog_client.post('/api/v1/catalog/uploads/'+items[0]['upload_id']+'/approve',headers=staff_headers,json=fields)
    assert response.status_code==200,response.text
    assert response.json()['data_origin']=='synthetic_demo'
    assert catalog_client.get('/api/v1/catalog/courses').json()['count']==0
    collection=FakeCollection()
    approved_index.publish_course('DEMO-PY-101',collection=collection,embedder=fake_embed)
    public=catalog_client.get('/api/v1/catalog/courses').json()['items']
    assert len(public)==1 and public[0]['data_origin']=='synthetic_demo'
    # An independently staged source can be rejected. A repeated setup preserves the audit decision.
    rejected=items[1]['upload_id']
    assert catalog_client.post('/api/v1/catalog/uploads/'+rejected+'/reject',headers=staff_headers,
             json={'reviewer':'TA Reviewer','reason':'Demonstrate source rejection'}).status_code==200
    again=stage(catalog_client,staff_headers)
    assert all(item['already_loaded'] for item in again)
    assert store.load_upload(rejected)['review_status']=='rejected'
    assert catalog_client.post('/api/v1/catalog/uploads/'+rejected+'/approve',headers=staff_headers,json=fields).status_code==409
    assert store.get_course('DEMO-PY-101')['status']=='published'


@pytest.mark.parametrize('suffix',['md','docx','pdf'])
def test_actual_demo_course_formats_upload(demo,catalog_client,staff_headers,suffix,example_pack):
    path=example_pack/'courses'/('demo-py-101.'+suffix)
    response=catalog_client.post('/api/v1/catalog/uploads',headers=staff_headers,files={'file':(path.name,path.read_bytes())})
    assert response.status_code==201,response.text
    record=store.load_upload(response.json()['upload_id'])
    assert record['chunks'] and record['data_origin']=='synthetic_demo'
    assert any('2400' in item['text'] for item in record['chunks'])


def test_examples_zip_contains_no_credentials_or_runtime_data(demo,catalog_client,staff_headers):
    response=catalog_client.get('/api/v1/demo/examples.zip',headers=staff_headers)
    assert response.status_code==200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert 'courses/demo-py-101.pdf' in archive.namelist()
        assert 'reference_materials/teaching_reference_pack.docx' in archive.namelist()
        assert not any('.env' in name or '.db' in name or name.startswith('/') for name in archive.namelist())


def test_examples_zip_excludes_unrelated_and_hidden_files(demo, catalog_client, staff_headers, example_pack):
    for name in ('.env', 'private.json', 'courses/.credentials.json', 'courses/runtime.sqlite3-wal'):
        (example_pack / name).write_text('must stay private', encoding='utf-8')
    response = catalog_client.get('/api/v1/demo/examples.zip', headers=staff_headers)
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert not any('must stay private' in archive.read(name).decode('utf-8', errors='ignore')
                       for name in archive.namelist())


@pytest.mark.parametrize('filename', ['../outside.md', 'C:/private.md', '/private.md', 'courses\\file.md', 'missing.md'])
def test_manifest_paths_are_validated_before_any_staging(demo, catalog_client, staff_headers, example_pack, filename):
    (example_pack / 'course_manifest.json').write_text(json.dumps({'courses': [
        {'file': 'courses/demo-py-101.md'}, {'file': filename}]}), encoding='utf-8')
    response = catalog_client.post('/api/v1/demo/stage-courses', headers=staff_headers, json={'reviewer': 'Reviewer'})
    assert response.status_code == 422
    assert catalog_client.get('/api/v1/catalog/uploads', headers=staff_headers).json()['count'] == 0


@pytest.mark.parametrize('endpoint,filename,content', [
    ('stage-courses', 'course_manifest.json', '{broken'),
    ('stage-courses', 'course_manifest.json', '{"courses": []}'),
    ('stage-tenders', 'tenders.json', '{"items": [null]}'),
    ('stage-tenders', 'tenders.json', '[]'),
])
def test_invalid_sample_pack_has_actionable_validation(demo, catalog_client, staff_headers, example_pack, endpoint, filename, content):
    (example_pack / filename).write_text(content, encoding='utf-8')
    response = catalog_client.post('/api/v1/demo/' + endpoint, headers=staff_headers, json={'reviewer': 'Reviewer'})
    assert response.status_code == 422


def test_symlink_sample_is_refused(demo, catalog_client, staff_headers, example_pack, tmp_path):
    secret = tmp_path / 'outside.md'
    secret.write_text('private content', encoding='utf-8')
    link = example_pack / 'courses/linked.md'
    try:
        link.symlink_to(secret)
    except OSError:
        pytest.skip('Creating symlinks requires OS privileges')
    response = catalog_client.get('/api/v1/demo/examples.zip', headers=staff_headers)
    assert response.status_code == 422
    assert b'private content' not in response.content


@pytest.mark.parametrize('variable', ['REX_CATALOG_DATA_DIR', 'RFP_DATA_DIR', 'REX_CURRICULUM_DIR'])
@pytest.mark.parametrize('relative', ['', 'nested/academy', '..'])
def test_demo_rejects_any_academy_storage_overlap(demo, monkeypatch, tmp_path, variable, relative):
    monkeypatch.setenv(variable, str((tmp_path / 'demo' / relative).resolve()))
    for resolver in (runtime.catalog_dir, runtime.rfp_dir, runtime.curriculum_dir):
        with pytest.raises(RuntimeError, match='must not overlap'):
            resolver()


def test_demo_rejects_academy_database_url_inside_demo_root(demo, monkeypatch, tmp_path):
    monkeypatch.setenv('RFP_DATABASE_URL', f'sqlite:///{tmp_path / "demo/rfp/rfp.db"}')
    with pytest.raises(RuntimeError, match='academy RFP database'):
        runtime.rfp_dir()


def test_faris_routes_use_host_auth_and_keep_dashboard_alerts_local(demo,catalog_client,staff_headers,monkeypatch):
    from backend.app.agents.rfp import notifications
    def forbidden(*a,**k): raise AssertionError('Must not send external messages')
    monkeypatch.setattr(notifications,'send_email_digest',forbidden)
    monkeypatch.setattr(notifications,'post_webhook',forbidden)
    assert catalog_client.get('/api/v1/rfp/opportunities').status_code==401
    response=catalog_client.post('/api/v1/demo/stage-tenders',headers=staff_headers)
    assert response.status_code==200,response.text
    result=response.json()
    assert result['new']==2 and result['rejected']==2
    items=catalog_client.get('/api/v1/rfp/opportunities',headers=staff_headers).json()['items']
    assert len(items)==2 and {x['country_code'] for x in items}=={'EG','SA'}
    assert all(x['data_mode']=='synthetic_demo' for x in items)
    assert catalog_client.get('/api/v1/rfp/notifications',headers=staff_headers).json()['count']>0
    oid=items[0]['id']
    assert catalog_client.patch(f'/api/v1/rfp/opportunities/{oid}',headers=staff_headers,
              json={'workflow_status':'reviewing','reviewer_notes':'TA checked'}).status_code==200
    assert catalog_client.patch(f'/api/v1/rfp/opportunities/{oid}',headers=staff_headers,
              json={'workflow_status':'invented'}).status_code==422
    activity=catalog_client.get('/api/v1/rfp/activity',headers=staff_headers)
    assert activity.status_code==200 and len(activity.json()['rejected'])==2
    assert catalog_client.post('/api/v1/rfp/run',headers=staff_headers,json={}).status_code==409
    again=catalog_client.post('/api/v1/demo/stage-tenders',headers=staff_headers).json()
    assert again['new']==0


def test_academy_rfp_refuses_synthetic_fixture(catalog_client,staff_headers,example_pack):
    data=json.loads((example_pack/'tenders.json').read_text())
    assert catalog_client.post('/api/v1/rfp/import',headers=staff_headers,json={'items':data['items']}).status_code==409


@pytest.mark.parametrize('configuration', ['', 'relative/path', '/missing/external/sample-pack'])
def test_missing_external_pack_returns_actionable_error(demo, catalog_client, staff_headers, monkeypatch, configuration):
    monkeypatch.setenv('REX_DEMO_EXAMPLES_DIR', configuration)
    for method, url in [('post', '/stage-courses'), ('post', '/stage-tenders'), ('get', '/examples.zip')]:
        kwargs = {'json': {'reviewer': 'Test Reviewer'}} if url == '/stage-courses' else {}
        result = getattr(catalog_client, method)('/api/v1/demo' + url, headers=staff_headers, **kwargs)
        assert result.status_code == 503
        assert 'REX_DEMO_EXAMPLES_DIR' in result.json()['detail']


def test_catalog_faq_validation():
    with pytest.raises(ValueError):
        store.CourseApproval(course_id='XX',title='Test course',description='A long enough description',reviewer='TA',faqs=[{'question':'','answer':'x'}])


def test_demo_refuses_same_physical_catalog_directory(demo, monkeypatch, tmp_path):
    monkeypatch.setenv("REX_CATALOG_DATA_DIR", str(tmp_path / "demo/catalog"))
    with pytest.raises(RuntimeError, match="must be different"):
        runtime.catalog_dir()


def test_copied_foreign_origin_record_is_not_public(published_course):
    course, _ = published_course
    with store.catalog_transaction() as con:
        row=con.execute("SELECT record_json FROM courses WHERE course_id=?", (course["course_id"],)).fetchone()
        payload=json.loads(row["record_json"])
        payload["data_origin"]="synthetic_demo"
        con.execute("UPDATE courses SET record_json=? WHERE course_id=?", (json.dumps(payload),course["course_id"]))
    assert store.list_courses()==[]
    assert not store.is_published(course["course_id"],course["source_sha256"])
    with pytest.raises(HTTPException):
        store.get_course(course["course_id"])
